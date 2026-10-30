import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from siget_desktop.ui.reportes_excel import guardar_excel_reportes


class ReportesExcelTests(unittest.TestCase):

    def test_generar_excel_reportes(self):
        datos = {
            "total_tickets_resueltos": 10,
            "tickets_evaluados_sla": 10,
            "tickets_dentro_sla": 8,
            "tickets_fuera_sla": 2,
            "tickets_sin_sla": 0,
            "porcentaje_cumplimiento_sla": 80.0,
            "tiempo_promedio_resolucion_horas": 3.5,
            "fallas_recurrentes": [
                {
                    "falla": "Falla de memoria RAM",
                    "cantidad": 3,
                }
            ],
            "componentes_reemplazados": [
                {
                    "id_repuesto": 1,
                    "repuesto": "Disco SSD",
                    "cantidad": 2,
                }
            ],
            "tasas_fallas_por_categoria": [
                {
                    "categoria": "Notebook",
                    "total_activos": 4,
                    "activos_con_fallas": 3,
                    "tasa_fallas_porcentaje": 75.0,
                }
            ],
        }

        with tempfile.TemporaryDirectory() as temporal:
            ruta = str(Path(temporal) / "reporte.xlsx")

            guardar_excel_reportes(
                ruta,
                datos,
                "2026-10-01",
                "2026-10-29",
            )

            archivo = Path(ruta)

            self.assertTrue(archivo.exists())
            self.assertGreater(archivo.stat().st_size, 0)

            libro = load_workbook(archivo, read_only=True)

            self.assertEqual(
                libro.sheetnames,
                [
                    "Resumen",
                    "Fallas recurrentes",
                    "Componentes reemplazados",
                    "Tasas por categoria",
                ],
            )

            # Comprobar resumen
            resumen = libro["Resumen"]

            self.assertEqual(
                resumen["A6"].value,
                "Total tickets resueltos",
            )
            self.assertEqual(resumen["B6"].value, 10)

            self.assertEqual(
                resumen["A10"].value,
                "Tickets sin SLA",
            )
            self.assertEqual(resumen["B10"].value, 0)

            # Comprobar fallas recurrentes
            fallas = libro["Fallas recurrentes"]

            self.assertEqual(fallas["A2"].value, "Falla de memoria RAM")
            self.assertEqual(fallas["B2"].value, 3)

            # Comprobar componentes reemplazados
            componentes = libro["Componentes reemplazados"]

            self.assertEqual(componentes["A2"].value, "Disco SSD")
            self.assertEqual(componentes["B2"].value, 2)

            # Comprobar tasas por categoría
            tasas = libro["Tasas por categoria"]

            self.assertEqual(tasas["A2"].value, "Notebook")
            self.assertEqual(tasas["B2"].value, 4)
            self.assertEqual(tasas["C2"].value, 3)
            self.assertEqual(tasas["D2"].value, 75.0)

            libro.close()


if __name__ == "__main__":
    unittest.main()
