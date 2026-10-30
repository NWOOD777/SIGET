from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

COLOR_AZUL = "172B4D"
COLOR_BORDE = "B8C4D2"
COLOR_FILA_ALTERNA = "F3F6FA"

BORDE_CELDA = Border(
    left=Side(style="thin", color=COLOR_BORDE),
    right=Side(style="thin", color=COLOR_BORDE),
    top=Side(style="thin", color=COLOR_BORDE),
    bottom=Side(style="thin", color=COLOR_BORDE),
)


def _estilizar_encabezados(
    hoja: Worksheet,
    fila_encabezado: int = 1,
) -> None:
    """Da formato a los encabezados sin crear filtros."""

    for celda in hoja[fila_encabezado]:
        celda.font = Font(
            bold=True,
            color="FFFFFF",
        )
        celda.fill = PatternFill(
            fill_type="solid",
            fgColor=COLOR_AZUL,
        )
        celda.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        celda.border = BORDE_CELDA

    hoja.row_dimensions[fila_encabezado].height = 30
    hoja.freeze_panes = f"A{fila_encabezado + 1}"


def _estilizar_datos(
    hoja: Worksheet,
    fila_inicio: int,
    cantidad_columnas: int,
) -> None:
    """Agrega bordes y alternancia de filas a los datos."""

    for fila in range(fila_inicio, hoja.max_row + 1):
        for columna in range(1, cantidad_columnas + 1):
            celda = hoja.cell(
                row=fila,
                column=columna,
            )

            celda.border = BORDE_CELDA
            celda.alignment = Alignment(
                horizontal=("left" if columna == 1 else "center"),
                vertical="center",
                wrap_text=True,
            )

            if (fila - fila_inicio) % 2 == 1:
                celda.fill = PatternFill(
                    fill_type="solid",
                    fgColor=COLOR_FILA_ALTERNA,
                )

        hoja.row_dimensions[fila].height = 27


def _ajustar_columnas(hoja: Worksheet) -> None:
    """Ajusta los anchos de las columnas."""

    for columna in hoja.columns:
        letra = get_column_letter(columna[0].column or 1)

        ancho = max(
            (
                len(str(celda.value)) if celda.value is not None else 0
                for celda in columna
            ),
            default=10,
        )

        hoja.column_dimensions[letra].width = min(
            max(ancho + 4, 15),
            55,
        )


def _crear_hoja(
    libro: Workbook,
    nombre: str,
    encabezados: list[str],
    filas: list[list[Any]],
) -> Worksheet:
    """Crea una hoja con tablas delimitadas, sin filtros."""

    hoja = libro.create_sheet(title=nombre)
    hoja.append(encabezados)

    for fila in filas:
        hoja.append(fila)

    _estilizar_encabezados(hoja)

    _estilizar_datos(
        hoja,
        fila_inicio=2,
        cantidad_columnas=len(encabezados),
    )

    _ajustar_columnas(hoja)

    return hoja


def guardar_excel_reportes(
    ruta: str,
    datos: dict[str, Any],
    fecha_inicio: str,
    fecha_fin: str,
) -> None:
    """Genera el informe Excel de SIGET sin filtros."""

    if Path(ruta).suffix.lower() != ".xlsx":
        raise ValueError("Seleccione un archivo con extensión .xlsx.")

    if not isinstance(datos, dict):
        raise ValueError("Los datos del informe no son válidos.")

    libro = Workbook()

    # ============================================================
    # HOJA 1: RESUMEN
    # ============================================================

    resumen = libro.active
    if not isinstance(resumen, Worksheet):
        raise RuntimeError("No se pudo crear la hoja Resumen del informe Excel.")

    resumen.title = "Resumen"

    resumen.append(["SIGET - INFORME DE GESTIÓN"])
    resumen.append(["Fecha de inicio", fecha_inicio])
    resumen.append(["Fecha de término", fecha_fin])
    resumen.append([])
    resumen.append(["Indicador", "Resultado"])

    metricas = [
        [
            "Total tickets resueltos",
            datos.get("total_tickets_resueltos", 0),
        ],
        [
            "Tickets evaluados SLA",
            datos.get("tickets_evaluados_sla", 0),
        ],
        [
            "Tickets dentro SLA",
            datos.get("tickets_dentro_sla", 0),
        ],
        [
            "Tickets fuera SLA",
            datos.get("tickets_fuera_sla", 0),
        ],
        [
            "Tickets sin SLA",
            datos.get("tickets_sin_sla", 0),
        ],
        [
            "Cumplimiento SLA (%)",
            datos.get("porcentaje_cumplimiento_sla", 0),
        ],
        [
            "Tiempo promedio de resolución (horas)",
            datos.get("tiempo_promedio_resolucion_horas", 0),
        ],
    ]

    for fila in metricas:
        resumen.append(fila)

    resumen["A1"].font = Font(
        size=16,
        bold=True,
        color=COLOR_AZUL,
    )

    resumen.row_dimensions[1].height = 30

    # Bordes para las fechas.
    for fila in (2, 3):
        for columna in (1, 2):
            celda = resumen.cell(fila, columna)
            celda.border = BORDE_CELDA
            celda.alignment = Alignment(
                vertical="center",
            )

    _estilizar_encabezados(
        resumen,
        fila_encabezado=5,
    )

    _estilizar_datos(
        resumen,
        fila_inicio=6,
        cantidad_columnas=2,
    )

    _ajustar_columnas(resumen)

    # ============================================================
    # HOJA 2: FALLAS RECURRENTES
    # ============================================================

    fallas = datos.get("fallas_recurrentes", []) or []

    _crear_hoja(
        libro,
        "Fallas recurrentes",
        ["Descripción de la falla", "Cantidad"],
        [
            [
                registro.get("falla", ""),
                registro.get("cantidad", 0),
            ]
            for registro in fallas
            if isinstance(registro, dict)
        ],
    )

    # ============================================================
    # HOJA 3: COMPONENTES REEMPLAZADOS
    # ============================================================

    componentes = datos.get("componentes_reemplazados", []) or []

    _crear_hoja(
        libro,
        "Componentes reemplazados",
        ["Repuesto", "Cantidad"],
        [
            [
                registro.get("repuesto", ""),
                registro.get("cantidad", 0),
            ]
            for registro in componentes
            if isinstance(registro, dict)
        ],
    )

    # ============================================================
    # HOJA 4: TASAS POR CATEGORÍA
    # ============================================================

    tasas = datos.get("tasas_fallas_por_categoria", []) or []

    _crear_hoja(
        libro,
        "Tasas por categoria",
        [
            "Categoría",
            "Total activos",
            "Activos con fallas",
            "Tasa de fallas (%)",
        ],
        [
            [
                registro.get("categoria", ""),
                registro.get("total_activos", 0),
                registro.get("activos_con_fallas", 0),
                registro.get("tasa_fallas_porcentaje", 0),
            ]
            for registro in tasas
            if isinstance(registro, dict)
        ],
    )

    # Garantizar que ninguna hoja tenga autofiltros.
    for hoja in libro.worksheets:
        hoja.auto_filter.ref = None

    # ============================================================
    # GUARDADO SEGURO
    # ============================================================

    archivo = Path(ruta)
    archivo.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    temporal = archivo.with_name(f".{archivo.stem}.tmp.xlsx")

    try:
        libro.save(temporal)
        temporal.replace(archivo)
    finally:
        if temporal.exists():
            temporal.unlink()
