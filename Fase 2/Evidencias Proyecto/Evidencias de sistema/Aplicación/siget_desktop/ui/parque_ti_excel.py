from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from siget_desktop.ui.reportes_excel import (
    COLOR_AZUL,
    _ajustar_columnas,
    _crear_hoja,
    _estilizar_datos,
    _estilizar_encabezados,
)


def _convertir_monto(valor: Any) -> Decimal:
    try:
        monto = Decimal(str(valor))
    except (InvalidOperation, TypeError, ValueError) as error:
        raise ValueError("El informe contiene un monto no válido.") from error
    if not monto.is_finite():
        raise ValueError("El informe contiene un monto no válido.")
    return monto


def _formatear_fecha_chilena(fecha: str) -> str:
    return datetime.strptime(fecha, "%Y-%m-%d").strftime("%d/%m/%Y")


def guardar_excel_parque_ti(
    ruta: str,
    datos: dict[str, Any],
    fecha_inicio: str,
    fecha_fin: str,
    departamento: str = "Todos los departamentos",
) -> None:
    """Exporta resumen, distribución y activos individuales a Excel."""
    archivo = Path(ruta)
    if archivo.suffix.lower() != ".xlsx":
        raise ValueError("Seleccione un archivo con extensión .xlsx.")
    if not isinstance(datos, dict):
        raise ValueError("Los datos del reporte no son válidos.")
    if datos.get("sin_resultados") or not datos.get("total_activos", 0):
        raise ValueError("No se encontraron activos para los criterios seleccionados")

    total_activos = int(datos["total_activos"])
    valorizacion = _convertir_monto(datos.get("valor_total_adquisicion", "0.00"))
    incluir_costos = datos.get("incluir_costos") is True
    costos = _convertir_monto(datos.get("costos_repuestos", "0.00"))
    distribucion = datos.get("distribucion_estados", [])
    detalle = datos.get("detalle_equipos", [])
    if not isinstance(distribucion, list):
        raise ValueError("La distribución de estados no es válida.")
    if not isinstance(detalle, list):
        raise ValueError("El detalle de equipos no es válido.")

    filas_estados = []
    for registro in distribucion:
        if not isinstance(registro, dict):
            raise ValueError("Existe un estado de activo no válido.")
        filas_estados.append(
            [str(registro.get("estado", "")), int(registro.get("cantidad", 0))]
        )

    filas_equipos = []
    for registro in detalle:
        if not isinstance(registro, dict):
            raise ValueError("Existe un equipo no válido.")
        filas_equipos.append(
            [
                str(registro.get("equipo", "")),
                str(registro.get("estado", "")),
                _convertir_monto(registro.get("valor_adquisicion", "0.00")),
            ]
        )

    libro = Workbook()
    resumen = libro.active
    if not isinstance(resumen, Worksheet):
        raise RuntimeError("No se pudo crear la hoja de resumen.")
    resumen.title = "Resumen parque TI"
    resumen.append(["SIGET - REPORTE CONSOLIDADO DEL PARQUE TI"])
    resumen.append(["Fecha de inicio", _formatear_fecha_chilena(fecha_inicio)])
    resumen.append(["Fecha de término", _formatear_fecha_chilena(fecha_fin)])
    resumen.append(["Departamento", departamento])
    resumen.append([])
    resumen.append(["Indicador", "Resultado"])
    resumen.append(["Total de activos", total_activos])
    resumen.append(["Valor total de adquisición", valorizacion])
    if incluir_costos:
        resumen.append(["Costos de repuestos", costos])
    resumen["A1"].font = Font(size=16, bold=True, color=COLOR_AZUL)
    resumen.row_dimensions[1].height = 32
    _estilizar_encabezados(resumen, fila_encabezado=6)
    _estilizar_datos(resumen, fila_inicio=7, cantidad_columnas=2)
    resumen["B8"].number_format = '"$"#,##0.00'
    if incluir_costos:
        resumen["B9"].number_format = '"$"#,##0.00'
    _ajustar_columnas(resumen)

    _crear_hoja(
        libro,
        "Distribucion estados",
        ["Estado del activo", "Cantidad"],
        filas_estados,
    )
    _crear_hoja(
        libro,
        "Detalle equipos",
        ["Equipo", "Estado", "Valor de adquisición"],
        filas_equipos,
    )
    hoja_detalle = libro["Detalle equipos"]
    hoja_detalle.column_dimensions["A"].width = 34
    hoja_detalle.column_dimensions["B"].width = 24
    hoja_detalle.column_dimensions["C"].width = 27
    for fila in range(2, hoja_detalle.max_row + 1):
        celda = hoja_detalle.cell(row=fila, column=3)
        celda.number_format = '"$"#,##0.00'
        celda.alignment = Alignment(horizontal="right", vertical="center")

    for hoja in libro.worksheets:
        hoja.auto_filter.ref = None

    archivo.parent.mkdir(parents=True, exist_ok=True)
    temporal = archivo.with_name(f".{archivo.stem}.tmp.xlsx")
    try:
        libro.save(temporal)
        temporal.replace(archivo)
    finally:
        if temporal.exists():
            temporal.unlink()
